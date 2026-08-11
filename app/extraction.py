"""Text extraction from resume files (PDF, DOCX, TXT) with graceful OCR fallback.

All extraction works on in-memory bytes (no disk writes) so it runs unchanged on
Vercel's read-only filesystem.
"""
from __future__ import annotations

import io
import re

from .config import get_settings

# Heuristic: if a PDF yields fewer than this many characters per page on average,
# it's probably a scanned/image PDF and needs OCR.
_MIN_CHARS_PER_PAGE = 40


class ExtractionResult:
    def __init__(self, text: str, warnings: list[str] | None = None) -> None:
        self.text = text or ""
        self.warnings = warnings or []

    @property
    def char_count(self) -> int:
        return len(self.text.strip())


def _clean(text: str) -> str:
    """Normalize whitespace while preserving line breaks (useful for name detection)."""
    text = text.replace("\x00", " ")
    # Collapse runs of spaces/tabs but keep newlines.
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf(data: bytes) -> ExtractionResult:
    warnings: list[str] = []
    try:
        import pdfplumber
    except ImportError:
        return ExtractionResult("", ["pdfplumber not installed; cannot read PDF."])

    pages_text: list[str] = []
    page_count = 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                pages_text.append(page.extract_text() or "")
    except Exception as exc:  # pragma: no cover - malformed pdf
        return ExtractionResult("", [f"Failed to parse PDF: {exc}"])

    text = _clean("\n".join(pages_text))
    avg = len(text) / page_count if page_count else 0

    if page_count and avg < _MIN_CHARS_PER_PAGE:
        ocr_text = _try_ocr(data)
        if ocr_text is not None:
            warnings.append("Low text density — used OCR fallback.")
            return ExtractionResult(_clean(ocr_text), warnings)
        warnings.append(
            "This looks like a scanned/image PDF, but OCR is unavailable "
            "(Tesseract not installed). Extracted text may be incomplete."
        )
    return ExtractionResult(text, warnings)


def _try_ocr(data: bytes) -> str | None:
    """Attempt OCR if both pytesseract and pdf2image (+ system tesseract) exist."""
    settings = get_settings()
    if not settings.has_ocr:
        return None
    try:  # pragma: no cover - OCR path not exercised in CI (no tesseract)
        import pdf2image
        import pytesseract

        images = pdf2image.convert_from_bytes(data)
        return "\n".join(pytesseract.image_to_string(img) for img in images)
    except Exception:
        return None


def _extract_docx(data: bytes) -> ExtractionResult:
    try:
        import docx  # python-docx
    except ImportError:
        return ExtractionResult("", ["python-docx not installed; cannot read DOCX."])
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        return ExtractionResult("", [f"Failed to parse DOCX: {exc}"])

    parts = [p.text for p in document.paragraphs]
    # Include table cell text (resumes often use tables for layout).
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return ExtractionResult(_clean("\n".join(parts)))


def _extract_txt(data: bytes) -> ExtractionResult:
    for encoding in ("utf-8", "latin-1"):
        try:
            return ExtractionResult(_clean(data.decode(encoding)))
        except UnicodeDecodeError:
            continue
    return ExtractionResult(_clean(data.decode("utf-8", errors="ignore")))


def extract_text(filename: str, data: bytes) -> ExtractionResult:
    """Dispatch on file extension. Returns text + any warnings."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _extract_pdf(data)
    if name.endswith(".docx"):
        return _extract_docx(data)
    if name.endswith((".txt", ".md")):
        return _extract_txt(data)
    if name.endswith(".doc"):
        return ExtractionResult(
            "", ["Legacy .doc format is not supported — please convert to .docx or PDF."]
        )
    # Last resort: try as text.
    return _extract_txt(data)


# --- Candidate name heuristic -------------------------------------------------

_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_NON_NAME = re.compile(r"(resume|curriculum vitae|cv|profile)", re.IGNORECASE)


def guess_candidate_name(text: str, filename: str) -> str | None:
    """Best-effort candidate name from the first meaningful line, else filename."""
    for line in text.splitlines():
        line = line.strip()
        if not line or _EMAIL_RE.search(line) or _NON_NAME.search(line):
            continue
        words = line.split()
        # A name is typically 2-4 capitalized-ish words, no digits.
        if 2 <= len(words) <= 4 and not any(ch.isdigit() for ch in line) and len(line) <= 48:
            return line.title() if line.isupper() else line
        break  # only inspect the first non-trivial line

    # Fall back to a cleaned-up filename.
    stem = re.sub(r"\.[^.]+$", "", filename or "").replace("_", " ").replace("-", " ")
    stem = _NON_NAME.sub("", stem).strip()
    return stem.title() if stem else None
